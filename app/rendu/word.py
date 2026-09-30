"""Export Word (.docx) de dépannage : mêmes données et mêmes textes que le PDF.

Le PDF reste le document de référence ; ce fichier sert quand un document modifiable
est exceptionnellement nécessaire. Le sommaire se met à jour à l'ouverture dans Word.
"""
import io
import os
import re
from datetime import date

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from ..moteur.analyse import IMPACT_ANOMALIES, LIBELLES_ANOMALIES, couverture_mdr
from ..moteur.outils import fmt_date, fmt_mois, fmt_nb, fmt_pct
from . import graphiques
from .pdf import (DESCRIPTIONS_MENACES, GLOSSAIRE, anomalies, appareils_tries, construire_plan, evolution,
                  groupes_tries, sans_heure, statut_taux, versions)

STATIQUE = os.path.normpath(os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "static"))
C1, C2, C3 = "00507A", "0099BA", "8DC21F"
ENCRE, ENCRE_DOUCE, BLANC = RGBColor(0x1F, 0x29, 0x33), RGBColor(0x5B, 0x67, 0x70), RGBColor(0xFF, 0xFF, 0xFF)
FOND_CLAIR, FOND_ZEBRE = "EEF4F8", "F6F9FB"
COULEURS = {
    "Critique": "C0392B", "Urgente": "C0392B", "Très élevé": "C0392B",
    "Élevé": "D35400", "Haute": "D35400", "Modéré": "B7791F", "Avertissement": "B7791F", "À surveiller": "B7791F",
    "Normale": "1E6FA8", "Faible": "3C8D2F", "Bon": "3C8D2F", "Normal": "3C8D2F", "Aucun": "3C8D2F",
    "Réalisée": "3C8D2F", "En attente": "5B6770", "À faire": "5B6770", "En cours": "5B6770", "Abandonnée": "5B6770",
}


def _rgb(h):
    return RGBColor.from_string(h)


def _d(iso):
    return date.fromisoformat(iso[:10]) if iso else None


# --------------------------------------------------------------------------- #
# Primitives XML
# --------------------------------------------------------------------------- #
def _ombrer(cell, couleur):
    shd = OxmlElement("w:shd")
    for k, v in (("w:val", "clear"), ("w:color", "auto"), ("w:fill", couleur)):
        shd.set(qn(k), v)
    cell._tc.get_or_add_tcPr().append(shd)


def _bordures_cellule(cell, **cotes):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for cote in ("top", "left", "bottom", "right"):
        val = cotes.get(cote, ("nil", 0, "FFFFFF"))
        el = OxmlElement(f"w:{cote}")
        el.set(qn("w:val"), val[0]); el.set(qn("w:sz"), str(val[1])); el.set(qn("w:color"), val[2])
        borders.append(el)
    tcPr.append(borders)


def _bordures_table(table):
    borders = OxmlElement("w:tblBorders")
    for cote in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{cote}")
        el.set(qn("w:val"), "single" if cote in ("top", "bottom", "insideH") else "nil")
        el.set(qn("w:sz"), "4"); el.set(qn("w:color"), "D5DDE3")
        borders.append(el)
    table._tbl.tblPr.append(borders)


def _marges(table, h=60, b=60, g=100, d=100):
    mar = OxmlElement("w:tblCellMar")
    for cote, v in (("top", h), ("bottom", b), ("left", g), ("right", d)):
        el = OxmlElement(f"w:{cote}")
        el.set(qn("w:w"), str(v)); el.set(qn("w:type"), "dxa")
        mar.append(el)
    table._tbl.tblPr.append(mar)


def _ligne_entete(row):
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    row._tr.get_or_add_trPr().append(el)


def _insecable(row):
    row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))


def _filet(paragraphe, couleur=C3, taille=12):
    pbdr = OxmlElement("w:pBdr")
    el = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", str(taille)), ("w:space", "3"), ("w:color", couleur)):
        el.set(qn(k), v)
    pbdr.append(el)
    paragraphe._p.get_or_add_pPr().append(pbdr)


def _champ(paragraphe, code, taille=8):
    run = paragraphe.add_run()
    for typ in ("begin", "instr", "separate", "texte", "end"):
        if typ == "instr":
            el = OxmlElement("w:instrText"); el.set(qn("xml:space"), "preserve"); el.text = f" {code} "
        elif typ == "texte":
            el = OxmlElement("w:t"); el.text = "1"
        else:
            el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), typ)
        run._r.append(el)
    run.font.size, run.font.color.rgb = Pt(taille), ENCRE_DOUCE


# --------------------------------------------------------------------------- #
# Document
# --------------------------------------------------------------------------- #
class Word:
    def __init__(self, d, c, presta):
        self.d, self.c, self.presta = d, c, presta
        self.doc = Document()
        self.numeros, self.titres, _ = construire_plan(d, c)
        self._styles()

    def _styles(self):
        st = self.doc.styles
        n = st["Normal"]
        n.font.name, n.font.size, n.font.color.rgb = "Calibri", Pt(10.5), ENCRE
        n.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        n.paragraph_format.space_after, n.paragraph_format.line_spacing = Pt(6), 1.12
        for nom, taille, couleur, avant, apres in (("Heading 1", 16, C1, 18, 8), ("Heading 2", 12.5, C2, 12, 5)):
            h = st[nom]
            rf = h.element.rPr.find(qn("w:rFonts"))
            if rf is not None:
                for att in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
                    rf.attrib.pop(qn(att), None)
                rf.set(qn("w:ascii"), "Calibri"); rf.set(qn("w:hAnsi"), "Calibri")
            h.font.size, h.font.bold, h.font.italic, h.font.color.rgb = Pt(taille), True, False, _rgb(couleur)
            h.paragraph_format.space_before, h.paragraph_format.space_after = Pt(avant), Pt(apres)
            h.paragraph_format.keep_with_next = True
        sec = self.doc.sections[0]
        sec.page_height, sec.page_width = Cm(29.7), Cm(21)
        sec.top_margin, sec.bottom_margin, sec.left_margin, sec.right_margin = Cm(2.2), Cm(2), Cm(2), Cm(2)
        # Demande à Word de mettre à jour le sommaire et les numéros de page à l'ouverture
        maj = OxmlElement("w:updateFields")
        maj.set(qn("w:val"), "true")
        self.doc.settings.element.append(maj)

    # ---- blocs ----------------------------------------------------------- #
    @staticmethod
    def _runs(p, texte, gras=False, taille=None, couleur=None, italique=False):
        for i, morceau in enumerate(re.split(r"\*\*", texte or "")):
            if morceau:
                r = p.add_run(morceau)
                r.bold, r.italic = gras or i % 2 == 1, italique
                if taille:
                    r.font.size = Pt(taille)
                if couleur:
                    r.font.color.rgb = couleur
        return p

    def para(self, texte="", gras=False, taille=None, couleur=None, italique=False, centre=False, apres=None):
        p = self._runs(self.doc.add_paragraph(), texte, gras, taille, couleur, italique)
        if centre:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if apres is not None:
            p.paragraph_format.space_after = Pt(apres)
        return p

    def texte(self, texte):
        """Texte relu : paragraphes séparés par une ligne vide."""
        for bloc in re.split(r"\n\s*\n", texte or ""):
            if bloc.strip():
                self.para(bloc.strip())

    def note(self, texte):
        return self.para(texte, italique=True, taille=8.5, couleur=ENCRE_DOUCE)

    def puce(self, texte):
        return self._runs(self.doc.add_paragraph(style="List Bullet"), texte)

    def h1(self, ident):
        h = self.doc.add_heading(f"{self.numeros[ident]} {self.titres[ident]}", level=1)
        _filet(h)

    def h2(self, ident):
        self.doc.add_heading(f"{self.numeros[ident]} {self.titres[ident]}", level=2)

    def saut_page(self):
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def image(self, contenu_png, largeur=16.5, legende=None):
        if not contenu_png:
            return
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(io.BytesIO(contenu_png), width=Cm(largeur))
        if legende:
            self.para(legende, italique=True, taille=8.5, couleur=ENCRE_DOUCE, centre=True, apres=10)

    def encadre(self, titre, lignes, couleur=C2, fond=FOND_CLAIR):
        t = self.doc.add_table(rows=1, cols=1)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        _marges(t, 100, 100, 180, 140)
        _insecable(t.rows[0])
        cell = t.cell(0, 0)
        _ombrer(cell, fond)
        _bordures_cellule(cell, left=("single", 24, couleur))
        p = cell.paragraphs[0]
        premier = True
        if titre:
            self._runs(p, titre, gras=True, couleur=_rgb(couleur))
            premier = False
        for ligne in lignes:
            q = p if premier else cell.add_paragraph()
            premier = False
            self._runs(q, ligne)
            q.paragraph_format.space_after = Pt(2)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def tableau(self, entetes, lignes, largeurs, statuts=(), centre=(), taille=9.5):
        t = self.doc.add_table(rows=1, cols=len(entetes))
        t.alignment, t.autofit = WD_TABLE_ALIGNMENT.CENTER, False
        _bordures_table(t)
        _marges(t)
        _ligne_entete(t.rows[0])
        for i, e in enumerate(entetes):
            cell = t.rows[0].cells[i]
            _ombrer(cell, C1)
            r = cell.paragraphs[0].add_run(e)
            r.bold, r.font.size, r.font.color.rgb = True, Pt(taille), BLANC
            if i in centre:
                cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        for n, ligne in enumerate(lignes):
            row = t.add_row()
            _insecable(row)
            for i, valeur in enumerate(ligne):
                cell = row.cells[i]
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                if n % 2:
                    _ombrer(cell, FOND_ZEBRE)
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(0)
                texte = "" if valeur is None else str(valeur)
                if i in statuts and texte in COULEURS:
                    r = p.add_run("● " + texte)
                    r.bold, r.font.size, r.font.color.rgb = True, Pt(taille), _rgb(COULEURS[texte])
                else:
                    self._runs(p, texte, taille=taille)
                if i in centre:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for row in t.rows:
            for i, w in enumerate(largeurs):
                row.cells[i].width = Cm(w)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)
        return t

    def tuiles(self, elements):
        t = self.doc.add_table(rows=1, cols=len(elements))
        t.alignment, t.autofit = WD_TABLE_ALIGNMENT.CENTER, False
        _marges(t, 110, 110, 90, 90)
        for i, (valeur, libelle, statut) in enumerate(elements):
            cell = t.cell(0, i)
            cell.width = Cm(17 / len(elements))
            _ombrer(cell, FOND_CLAIR)
            couleur = COULEURS.get(statut, C1)
            _bordures_cellule(cell, top=("single", 24, couleur), left=("single", 18, "FFFFFF"), right=("single", 18, "FFFFFF"))
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(str(valeur))
            r.bold, r.font.size, r.font.color.rgb = True, Pt(20), _rgb(couleur)
            p.paragraph_format.space_after = Pt(0)
            q = cell.add_paragraph()
            q.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = q.add_run(libelle)
            r.font.size, r.font.color.rgb = Pt(8.5), ENCRE_DOUCE
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def en_attente(self, nom):
        self.encadre("Données en attente", [
            f"L'export Kaspersky Security Center « {nom} » n'a pas encore été transmis pour cette période.",
            "Cette section sera complétée dès réception du fichier."], couleur="5B6770", fond="F3F4F6")

    # ---- en-tête, pied, couverture, sommaire ------------------------------ #
    def cadre(self, date_rapport):
        d, c, presta = self.d, self.c, self.presta
        logo = os.path.join(STATIQUE, "logo.png")
        mois = fmt_mois(d["annee"], d["mois"])
        sec = self.doc.sections[0]
        sec.different_first_page_header_footer = True
        t = sec.header.add_table(rows=1, cols=2, width=Cm(17))
        t.cell(0, 0).paragraphs[0].add_run().add_picture(logo, height=Cm(0.95))
        p = t.cell(0, 1).paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r = p.add_run(f"Rapport mensuel de sécurité — {d['client']}\n{mois}")
        r.font.size, r.font.color.rgb = Pt(8.5), ENCRE_DOUCE
        _filet(sec.header.add_paragraph(), taille=8)
        f = sec.footer.paragraphs[0]
        f.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = f.add_run(f"Confidentiel — {presta['nom']} / {presta['signataire']}     ·     Page ")
        r.font.size, r.font.color.rgb = Pt(8), ENCRE_DOUCE
        _champ(f, "PAGE")
        r = f.add_run(" / ")
        r.font.size, r.font.color.rgb = Pt(8), ENCRE_DOUCE
        _champ(f, "NUMPAGES")

        p = self.doc.paragraphs[0] if self.doc.paragraphs else self.doc.add_paragraph()
        p.add_run().add_picture(logo, height=Cm(2.4))
        p.paragraph_format.space_after = Pt(60)
        bande = self.doc.add_table(rows=1, cols=1)
        _marges(bande, 360, 360, 400, 400)
        cell = bande.cell(0, 0)
        _ombrer(cell, C1)
        _bordures_cellule(cell, bottom=("single", 36, C3))
        for i, (texte, taille, couleur) in enumerate((("RAPPORT MENSUEL", 28, BLANC), ("DE SÉCURITÉ", 28, _rgb(C3)),
                                                       (f"{d['client']}  ·  {mois}", 15, BLANC))):
            q = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
            r = q.add_run(texte)
            r.bold, r.font.size, r.font.color.rgb = i < 2, Pt(taille), couleur
            q.paragraph_format.space_after = Pt(0)
        self.para(apres=26)
        sources = " · ".join(filter(None, ["Kaspersky MDR" if d["avec_mdr"] else "",
                                            "Kaspersky Security Center" if d["exports"] else ""])) or "—"
        infos = [("Client", d["client"]), ("Prestataire", presta["nom_complet"]),
                 ("Période couverte", f"Du {fmt_date(_d(d['debut']))} au {fmt_date(_d(d['fin_incluse']))}"),
                 ("Date du rapport", fmt_date(date_rapport)), ("Rédacteur", presta["signataire"]), ("Sources", sources)]
        t = self.doc.add_table(rows=0, cols=2)
        t.autofit = False
        _bordures_table(t)
        _marges(t, 80, 80, 120, 120)
        for libelle, valeur in infos:
            a, b = t.add_row().cells
            a.width, b.width = Cm(4.5), Cm(12.5)
            _ombrer(a, FOND_CLAIR)
            r = a.paragraphs[0].add_run(libelle)
            r.bold, r.font.color.rgb = True, _rgb(C1)
            b.paragraphs[0].add_run(valeur)
        self.para(apres=40)
        niveau = c["niveau_risque"]
        self.para("Niveau de risque global", taille=10, couleur=ENCRE_DOUCE, centre=True, apres=0)
        self.para(niveau.upper(), gras=True, taille=22, couleur=_rgb(COULEURS[niveau]), centre=True, apres=40)
        self.note("Document confidentiel, destiné exclusivement au client mentionné ci-dessus.").alignment = WD_ALIGN_PARAGRAPH.CENTER
        self.saut_page()

        h = self.para("Sommaire", gras=True, taille=16, couleur=_rgb(C1), apres=8)
        _filet(h)
        run = self.doc.add_paragraph().add_run()
        for typ, texte in (("begin", None), ("instr", 'TOC \\o "1-2" \\h \\z \\u'), ("separate", None),
                           ("t", "Clic droit → « Mettre à jour les champs » pour afficher le sommaire."), ("end", None)):
            if typ == "instr":
                el = OxmlElement("w:instrText"); el.set(qn("xml:space"), "preserve"); el.text = texte
            elif typ == "t":
                el = OxmlElement("w:t"); el.text = texte
            else:
                el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), typ)
            run._r.append(el)
        self.saut_page()

    # ---- sections ---------------------------------------------------------- #
    def synthese(self):
        d, c = self.d, self.c
        p, m, v, mdr = d["protection"], d["menaces"], d["vulnerabilites"], d["mdr"]
        self.h1("synthese")
        self.texte(c["synthese"])
        niveau = c["niveau_risque"]
        lignes = [f"• {x}" for x in c["motifs_risque"]] or ["• Aucun point bloquant relevé sur la période."]
        if d["risque"]["partiel"]:
            lignes.append("Évaluation partielle : les exports du parc (KSC) ne sont pas tous disponibles.")
        self.encadre(f"Niveau de risque global : {niveau.upper()}", lignes, couleur=COULEURS[niveau])
        tuiles = []
        if mdr:
            tuiles += [(fmt_nb(mdr["max"]), "postes max. supervisés (MDR)", None if mdr["max"] else "Critique"),
                       (len(mdr["incidents"]), "incidents MDR", "Critique" if mdr["incidents"] else "Bon")]
        if p:
            tuiles.append((f"{p['critique']}/{p['total']}", "appareils en état critique", statut_taux(p["critique"], p["total"])))
        if m:
            tuiles.append((fmt_nb(m["detections"]), "menaces détectées", None))
        if v:
            tuiles.append((v["critiques"], "vulnérabilités critiques", "Critique" if v["critiques"] else "Bon"))
        if tuiles:
            self.tuiles(tuiles[:5])
        ev = lambda cle: evolution(d, cle)  # noqa: E731
        lignes = []
        if mdr:
            cmdr = couverture_mdr(mdr, p)
            lignes += [["Postes supervisés par le MDR (moyenne jours ouvrés)",
                        str(round(mdr["moyenne_ouvres"] or mdr["moyenne"], 1)).replace(".", ","),
                        cmdr["statut"], ev("postes_mdr_moyenne")]]
            if cmdr["taux"] is not None:
                lignes += [["Couverture MDR du parc (maximum supervisé / appareils administrés)",
                            f"{cmdr['supervises']} / {cmdr['administres']} ({fmt_pct(cmdr['supervises'], cmdr['administres'])})",
                            cmdr["statut"], "—"]]
            lignes += [
                       ["Incidents de sécurité MDR", str(len(mdr["incidents"])),
                        "Critique" if mdr["incidents"] else "Aucun", ev("incidents_mdr")]]
        if p:
            sans = p["anomalies"].get("protection_desactivee", 0) + p["anomalies"].get("non_installe", 0)
            lignes += [["Appareils administrés (KSC)", str(p["total"]), "Normal", ev("appareils_administres")],
                       ["Appareils en état critique", f"{p['critique']} ({fmt_pct(p['critique'], p['total'])})",
                        statut_taux(p["critique"], p["total"]), ev("appareils_critiques")],
                       ["Appareils sans protection active", str(sans), "Critique" if sans else "Bon", "—"]]
        else:
            lignes.append(["État de la protection du parc", "—", "En attente", "—"])
        if m:
            lignes += [["Menaces détectées", f"{fmt_nb(m['detections'])} événements",
                        "Bon" if not m["non_neutralisees"] else "À surveiller", ev("detections")],
                       ["Appareils exposés à une menace", str(m["appareils_touches"]),
                        "Modéré" if m["appareils_touches"] else "Bon", ev("appareils_touches")]]
        else:
            lignes.append(["Menaces détectées", "—", "En attente", "—"])
        lignes.append(["Vulnérabilités (dont critiques)", f"{v['total']} ({v['critiques']})",
                       "Critique" if v["critiques"] else "Bon", ev("vulnerabilites_critiques")] if v
                      else ["Vulnérabilités logicielles", "—", "En attente", "—"])
        self.tableau(["Indicateur", "Valeur", "Statut", "Évolution*"], lignes, [7.5, 3.5, 3, 3], statuts=(2,), centre=(1, 3))
        self.note("* Évolution par rapport au mois précédent." if d["precedents"]
                  else "* Premier rapport de la série : l'évolution sera calculée à partir du mois prochain.")
        prio = [a for a in c["actions"] if a["priorite"] == "Urgente"][:3] or c["actions"][:3]
        if prio:
            self.para("**Priorités du mois**", apres=3)
            for a in prio:
                self.puce(f"{a['action']} — {a['responsable']}")
        if c.get("suggestion_mdr"):
            self.encadre("Recommandation : service Kaspersky MDR", [c["suggestion_mdr"]])

    def mdr(self, g):
        d, c, mdr, p = self.d, self.c, self.d["mdr"], self.d["protection"]
        if not d["avec_mdr"]:
            return
        self.h1("mdr")
        self.para("Le service Kaspersky MDR collecte en continu la télémétrie des postes équipés, la corrèle et fait "
                  "analyser les activités suspectes par des analystes SOC 24h/24 et 7j/7.")
        self.h2("mdr-couverture")
        if not mdr["tenant_trouve"] or not mdr["max"]:
            self.encadre("Aucune télémétrie reçue", [f"Aucun poste de {d['client']} n'a transmis de télémétrie au service "
                                                     "MDR sur la période."], couleur=COULEURS["Critique"])
            return
        self.texte(c["commentaires"]["mdr"])
        self.image(g.get("mdr"), legende="Nombre de postes ayant transmis de la télémétrie, par jour")
        lignes = [["Moyenne des jours ouvrés", str(round(mdr["moyenne_ouvres"], 1)).replace(".", ",")],
                  ["Maximum / minimum journalier", f"{mdr['max']} / {mdr['min']}"],
                  ["Jours couverts par les données", f"{len(mdr['jours'])} sur {len(mdr['jours']) + len(mdr['jours_manquants'])}"]]
        if p:
            lignes.append(["Taux de couverture MDR (max. / parc)", fmt_pct(mdr["max"], p["total"])])
        self.tableau(["Indicateur", "Valeur"], lignes, [11, 6], centre=(1,))
        self.h2("mdr-incidents")
        if not mdr["incidents"]:
            self.encadre("Aucun incident confirmé", ["Aucun incident de sécurité n'a été créé ou traité par les analystes "
                                                     "MDR pour ce client sur la période."], couleur=COULEURS["Bon"])
        else:
            self.tableau(["N°", "Incident", "Créé le", "Priorité", "Statut", "Résolution"],
                         [[i["numero"], i["nom"], i["cree"], i["priorite"], i["statut"], i["resolution"]]
                          for i in mdr["incidents"]], [1.5, 5.5, 2.8, 2, 2.4, 2.8])

    def protection(self, g):
        p = self.d["protection"]
        self.h1("protection")
        if not p:
            return self.en_attente("Rapport sur l'état de la protection")
        self.texte(self.c["commentaires"]["protection"])
        self.image(g.get("parc"), 16, "Répartition des appareils par état de protection")
        self.h2("anomalies")
        if p.get("diagnostic"):
            self.para("Chaque appareil en anomalie est accompagné, dans la console, de la raison de son état. Ces raisons sont "
                      "regroupées ci-dessous avec leur cause probable et l'action recommandée.")
            self.tableau(["Raison signalée", "Appareils", "Impact", "Recommandation"],
                         [[f"**{l['libelle']}** — cause probable : {l['cause']}",
                           f"{l['appareils']} ({l['critiques']} crit." + (f", {l['avertissements']} avert." if l["avertissements"] else "") + ")",
                           l["impact"], f"{l['recommandation']} Responsable : {l['responsable']}."]
                          for l in p["diagnostic"]], [5.2, 2.2, 2.3, 7.3], statuts=(2,), centre=(1,), taille=8.5)
            if p.get("raisons_inconnues"):
                self.para("Raisons non répertoriées, à examiner dans la console : "
                          + " ; ".join(f"« {t} » ({n})" for t, n in p["raisons_inconnues"]) + ".", italique=True, taille=8.5)
        else:
            self.tableau(["Anomalie", "Appareils", "Impact"],
                         [[LIBELLES_ANOMALIES[k], n, IMPACT_ANOMALIES[k]] for k, n in p["anomalies"].items()],
                         [10.5, 2.5, 4], statuts=(2,), centre=(1,))
        if p.get("classement"):
            self.h2("classement")
            self.para("Appareils regroupés par état puis par combinaison de raisons : un même groupe se traite par une même action.")
            self.tableau(["État", "Raison(s)", "Appareils", "Appareils concernés"],
                         [[g["etat"], g["libelle"], g["appareils"],
                           ", ".join(g["noms"][:8]) + (f" … et {len(g['noms']) - 8} autres" if len(g["noms"]) > 8 else "")]
                          for g in p["classement"]], [2.5, 6, 2, 6.5], statuts=(0,), centre=(2,), taille=8.5)
        if p["serveurs"]:
            self.h2("serveurs")
            self.tableau(["Serveur", "Système", "État", "Anomalies"],
                         [[s["appareil"], s["os"].replace("Microsoft ", ""), s["etat"], anomalies(s["anomalies"])]
                          for s in p["serveurs"]], [3.8, 3.7, 2.5, 7], statuts=(2,), taille=9)
        if p["hors_ligne"]:
            self.h2("deconnectes")
            self.tableau(["Appareil", "Service", "Dernière connexion", "Depuis"],
                         [[a["appareil"], a["groupe"], sans_heure(a["derniere_connexion"]), f"{a['jours']} jours"]
                          for a in p["hors_ligne"][:15]], [4.5, 4.5, 5, 3], centre=(3,), taille=9)
        if p["os_obsoletes"]:
            self.h2("systemes")
            self.tableau(["Système", "Appareils", "Support Microsoft", "Statut"],
                         [[o["os"], o["appareils"], ("Support terminé depuis le " if o["expire"] else "Fin de support le ")
                           + fmt_date(_d(o["fin_support"])), "Critique" if o["expire"] else "Avertissement"]
                          for o in p["os_obsoletes"]], [4.5, 2.2, 7, 3.3], statuts=(3,), centre=(1,))
        self.h2("services")
        self.tableau(["Service / groupe", "Critique", "Avertissement", "Total en anomalie"],
                     [[grp, cpt.get("Critique", 0), cpt.get("Avertissement", 0), sum(cpt.values())]
                      for grp, cpt in groupes_tries(p["groupes"])], [7, 3, 3, 4], centre=(1, 2, 3))

    def menaces(self, g):
        m = self.d["menaces"]
        self.h1("menaces")
        if not m:
            return self.en_attente("Rapport sur les menaces")
        self.texte(self.c["commentaires"]["menaces"])
        if not m["detections"]:
            return
        self.h2("typologie")
        self.image(g.get("categories"), 15, "Nombre de détections par catégorie de menace")
        self.tableau(["Catégorie", "Description", "Détections", "Exemples"],
                     [[cat, DESCRIPTIONS_MENACES.get(cat, ""), n,
                       ", ".join(m["domaines_phishing"][:4] if cat == "Phishing" else m["exemples"].get(cat, [])) or "—"]
                      for cat, n in m["categories"].items()], [3.2, 5.6, 2.2, 6], centre=(2,), taille=8.5)
        self.h2("chronologie")
        self.image(g.get("menaces_jour"), legende="Nombre de détections par jour")
        self.h2("exposes")
        self.tableau(["Appareil", "Détections", "Menace principale"], [[a, n, cat] for a, n, cat in m["top_appareils"]],
                     [6, 3, 8], centre=(1,))

    def vulnerabilites(self, g):
        v = self.d["vulnerabilites"]
        self.h1("vulnerabilites")
        if not v:
            return self.en_attente("Rapport sur les vulnérabilités")
        self.texte(self.c["commentaires"]["vulnerabilites"])
        self.image(g.get("vulns"), legende="Vulnérabilités par application et par gravité")
        self.h2("applications")
        self.tableau(["Application", "Vuln.", "Critiques", "Appareils*", "Versions détectées", "Sévérité"],
                     [[a, i["total"], i["critiques"], i["appareils"], versions(i["versions"]),
                       "Critique" if i["critiques"] else "Élevé" if i["elevees"] else "Modéré"]
                      for a, i in v["applications"][:12]], [4, 1.5, 1.8, 2, 4.9, 2.8], statuts=(5,), centre=(1, 2, 3), taille=8.5)
        self.note("* Nombre maximal d'appareils exposés à une même vulnérabilité de l'application.")

    def actions(self):
        c = self.c
        self.h1("actions")
        if c.get("suivi"):
            self.h2("actions-suivi")
            self.tableau(["Action", "Priorité initiale", "Responsable", "Avancement"],
                         [[f"**{a['action']}**" + (f"\n{a['commentaire']}" if a.get("commentaire") else ""),
                           a["priorite"], a["responsable"], a["statut"]] for a in c["suivi"]],
                         [8.5, 2.8, 3, 2.7], statuts=(1, 3), taille=9)
            self.h2("actions-mois")
        if not c["actions"]:
            return self.para("Aucune action corrective n'est requise ce mois-ci.")
        self.tableau(["#", "Action et justification", "Priorité", "Responsable"],
                     [[i + 1, f"**{a['action']}**" + (f"\n{a['pourquoi']}" if a.get("pourquoi") else ""),
                       a["priorite"], a["responsable"]] for i, a in enumerate(c["actions"])],
                     [0.8, 10.7, 2.5, 3], statuts=(2,), centre=(0,), taille=9)

    def conclusion(self):
        self.h1("conclusion")
        self.texte(self.c["conclusion"])
        self.para(apres=14)
        self.para(self.presta["signataire"], gras=True, couleur=_rgb(C1), apres=0)
        self.para(self.presta["nom_complet"], couleur=ENCRE_DOUCE)

    def annexes(self):
        d, p = self.d, self.d["protection"]
        self.saut_page()
        if p:
            h = self.doc.add_heading(f"{self.numeros['annexe-appareils']} {self.titres['annexe-appareils']}", level=1)
            _filet(h)
            self.tableau(["Appareil", "Service", "État", "Anomalies"],
                         [[a["appareil"], a["groupe"], a["etat"], anomalies(a["anomalies"]) or a["raison"]]
                          for a in appareils_tries(p["appareils"])], [3.8, 3.4, 2.6, 7.2], statuts=(2,), taille=8)
        h = self.doc.add_heading(f"{self.numeros['annexe-sources']} {self.titres['annexe-sources']}", level=1)
        _filet(h)
        lignes = [["Kaspersky MDR (hebdomadaire)", s["fichier"], f"{s['debut']} → {s['fin']}"]
                  for s in (d["mdr"]["semaines"] if d["mdr"] else [])]
        for typ, nom in (("protection", "KSC — État de la protection"), ("menaces", "KSC — Menaces"),
                         ("vulnerabilites", "KSC — Vulnérabilités")):
            e = d["exports"].get(typ)
            lignes.append([nom, e["fichier"] if e else "Non transmis", (e.get("genere_le") or "—") if e else "—"])
        self.tableau(["Source", "Fichier", "Période / date"], lignes, [5, 8, 4], taille=8.5)
        for a in d["avertissements"]:
            self.puce(a)
        h = self.doc.add_heading(f"{self.numeros['annexe-glossaire']} {self.titres['annexe-glossaire']}", level=1)
        _filet(h)
        self.tableau(["Terme", "Définition"], [[f"**{t}**", df] for t, df in GLOSSAIRE], [3, 14], taille=9)


def docx_rapport(d, contenu, prestataire, date_rapport=None):
    """Construit le .docx et renvoie son contenu (octets)."""
    w = Word(d, contenu, prestataire)
    g = graphiques.tous(d, "png")
    w.cadre(date_rapport or date.today())
    w.synthese()
    w.mdr(g)
    w.protection(g)
    w.menaces(g)
    w.vulnerabilites(g)
    w.actions()
    w.conclusion()
    w.annexes()
    tampon = io.BytesIO()
    w.doc.save(tampon)
    return tampon.getvalue()
